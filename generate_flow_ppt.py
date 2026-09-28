"""
Bhairava Anugraha Super App - Flow Presentation Generator
Generates a comprehensive, professional, 16:9 widescreen PowerPoint presentation (PPTX)
detailing the architecture, user journeys, spiritual progression engine, and administrative flows.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# ---------------------------------------------------------
# CONSTANTS & THEME PALETTE (Sacred Dark-Gold Temple Codex)
# ---------------------------------------------------------
BG_COLOR = RGBColor(12, 10, 8)            # Deep Temple Void (#0c0a08)
CARD_BG = RGBColor(22, 18, 14)            # Dark Card Charcoal (#16120e)
CARD_BORDER = RGBColor(96, 73, 34)        # Muted Antique Gold Border (#604922)
CARD_BG_LIGHT = RGBColor(32, 26, 20)      # Lighter Card Panel (#201a14)

GOLD_PRIMARY = RGBColor(212, 166, 74)     # Radiant Gold (#d4a64a)
GOLD_BRIGHT = RGBColor(245, 222, 150)     # Pale Champagne Gold (#f5de96)
GOLD_DARK = RGBColor(166, 122, 44)        # Deep Amber (#a67a2c)

CRIMSON = RGBColor(158, 30, 30)           # Temple Crimson (#9e1e1e)
CRIMSON_BG = RGBColor(46, 14, 14)         # Muted Crimson Backing (#2e0e0e)

TEXT_LIGHT = RGBColor(238, 230, 218)      # Soft Parchment (#eee6da)
TEXT_MUTED = RGBColor(168, 158, 146)      # Slate / Stone Grey (#a89e92)
TEXT_WHITE = RGBColor(255, 255, 255)

ACCENT_GREEN = RGBColor(46, 125, 50)      # Approved / Active Green
ACCENT_BLUE = RGBColor(30, 100, 160)      # Info Blue
ACCENT_AMBER = RGBColor(217, 119, 6)      # Pending Amber

FONT_HEADING = 'Georgia'
FONT_BODY = 'Calibri'
FONT_MONO = 'Consolas'

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(SCRIPT_DIR, 'static', 'images')


def set_slide_background(slide):
    """Sets a dark background color for the slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_header(slide, title_text, category_text="BHAIRAVA ANUGRAHA SUPER APP"):
    """Adds a standardized sacred dark-gold header with category badge."""
    # Category / Tagline
    tx_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_cat = tx_cat.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"✦ {category_text.upper()}"
    p_cat.font.name = FONT_BODY
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = GOLD_PRIMARY

    # Main Title
    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = GOLD_BRIGHT

    # Thin decorative divider line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = CARD_BORDER
    line.line.color.rgb = CARD_BORDER


def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    """Creates a styled card container."""
    card = slide.shapes.add_shape(shape, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
    else:
        card.line.fill.background()
    return card


def add_footer(slide, current_slide, total_slides=13):
    """Adds a discreet footer with page numbering."""
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = tx_box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"॥ सर्वं भैरव स्वरूपम् ॥   |   Bhairava Anugraha Super App Architecture Deck   |   Slide {current_slide} of {total_slides}"
    p.font.name = FONT_BODY
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_MUTED


def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE & COVER SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Outer decorative framed border
    outer_frame = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5), Inches(0.5), Inches(12.333), Inches(6.5)
    )
    outer_frame.fill.background()
    outer_frame.line.color.rgb = CARD_BORDER
    outer_frame.line.width = Pt(1.5)

    inner_frame = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3)
    )
    inner_frame.fill.background()
    inner_frame.line.color.rgb = GOLD_DARK
    inner_frame.line.width = Pt(0.75)

    # Optional Logo
    logo_path = os.path.join(IMG_DIR, 'logo.jpeg')
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(1.0), Inches(1.8), Inches(2.2), Inches(2.2))

    # Sanskrit Invocation
    inv_box = slide1.shapes.add_textbox(Inches(3.5), Inches(1.4), Inches(8.8), Inches(0.6))
    p_inv = inv_box.text_frame.paragraphs[0]
    p_inv.text = "॥ ॐ श्री महाकाल भैरवाय नमः ॥"
    p_inv.font.name = FONT_HEADING
    p_inv.font.size = Pt(20)
    p_inv.font.color.rgb = GOLD_PRIMARY

    # Main Title
    t_box = slide1.shapes.add_textbox(Inches(3.5), Inches(2.0), Inches(8.8), Inches(1.6))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p1 = tf_t.paragraphs[0]
    p1.text = "BHAIRAVA ANUGRAHA SUPER APP"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = GOLD_BRIGHT

    p2 = tf_t.add_paragraph()
    p2.text = "Comprehensive Application Architecture, User Journeys & End-to-End Operational Flows"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_LIGHT

    # Overview pill cards at bottom of title
    pills = [
        ("Unified Sanctuary", "Landing, Ashtami & Media"),
        ("Jnāna Samvāda", "257+ Guruji Q&A Codex"),
        ("Bhairav Loka", "324 Sahasralinga Codex"),
        ("Sādhana Progression", "Gated 9-Stage Diksha Path"),
        ("Admin & Governance", "Sync, Chat & Analytics")
    ]
    pill_w = Inches(2.2)
    pill_gap = Inches(0.18)
    start_x = Inches(1.0)
    for i, (p_title, p_sub) in enumerate(pills):
        px = start_x + i * (pill_w + pill_gap)
        card = add_card(slide1, px, Inches(4.5), pill_w, Inches(1.5), bg_color=CARD_BG, border_color=CARD_BORDER)
        tx = slide1.shapes.add_textbox(px + Inches(0.1), Inches(4.6), pill_w - Inches(0.2), Inches(1.3))
        tf = tx.text_frame
        tf.word_wrap = True
        p_hdr = tf.paragraphs[0]
        p_hdr.text = p_title
        p_hdr.font.name = FONT_HEADING
        p_hdr.font.size = Pt(13)
        p_hdr.font.bold = True
        p_hdr.font.color.rgb = GOLD_PRIMARY

        p_desc = tf.add_paragraph()
        p_desc.text = p_sub
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED

    # Bottom Metadata
    meta_box = slide1.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.4))
    p_meta = meta_box.text_frame.paragraphs[0]
    p_meta.text = "Python Flask Architecture  •  Zero-Modification Sub-App Proxy  •  SQLAlchemy ORM  •  Temple Gold Codex Aesthetics"
    p_meta.font.name = FONT_MONO
    p_meta.font.size = Pt(10)
    p_meta.font.color.rgb = GOLD_DARK

    # Speaker Notes
    slide1.notes_slide.notes_text_frame.text = (
        "Welcome to the comprehensive Flow Presentation for the Bhairava Anugraha Super App.\n"
        "This deck walks through the complete end-to-end technical architecture, user personas, "
        "public discovery flows, gated spiritual progression engine, real-time guidance chat, "
        "and administrative governance."
    )

    # =========================================================================
    # SLIDE 2: EXECUTIVE VISION & PLATFORM OVERVIEW
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Executive Vision & System Unification", "Platform Overview")
    add_footer(slide2, 2)

    # Left Column: Strategic Vision & Context
    add_card(slide2, Inches(0.8), Inches(1.6), Inches(4.0), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_l = slide2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.6), Inches(4.7))
    tf_l = tx_l.text_frame
    tf_l.word_wrap = True
    
    p = tf_l.paragraphs[0]
    p.text = "The Sacred Vision"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD_PRIMARY

    p_body = tf_l.add_paragraph()
    p_body.text = (
        "The Bhairava Anugraha Super App unifies previously disconnected spiritual "
        "repositories, sacred texts, and darshan archives into a singular, high-performance web sanctuary.\n\n"
        "Core Objectives:\n"
        "• Unify bhairavaanugraha.com official offerings with the Jnāna Samvāda and Bhairav Loka Sahasralinga archives.\n"
        "• Enforce a Zero-Modification Guarantee across underlying upstream repositories.\n"
        "• Provide a structured, multi-tier Sādhana Diksha progression system with administrative gatekeeping.\n"
        "• Deliver an immersive dark-gold temple codex visual & acoustic aesthetic."
    )
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = TEXT_LIGHT

    # Right Column: 4 Core Pillars
    pillars_info = [
        ("1. Official Sanctuary Portal", "Route: /", "Faithful recreation of bhairavaanugraha.com with live Ashtami lunar countdowns, Guruji discourse videos, 'Life Within or Without' sacred book showcase, and WhatsApp/Telegram community links.", GOLD_PRIMARY),
        ("2. Jnāna Samvāda Codex", "Route: /jnana-samvada", "Curated archive of 257+ verified answers by Guruji categorized into 6 sacred Folios (Mantra, Puja, Mandala, Sādhana Experiences, Advanced, Women & Sādhana) with instant ⌘K fuzzy search.", GOLD_BRIGHT),
        ("3. Bhairav Loka Sahasralinga", "Route: /bhairav-loka", "Interactive celestial logarithmic spiral rendering 324 consecrated Shiva Lingams with darshan lightbox, post transcripts, and view modes (Spiral, Grid, Matrix, Cards).", GOLD_PRIMARY),
        ("4. Sādhana Progression Engine", "Routes: /sadhana-paddhati, /stage/<n>", "Hierarchical 9-stage spiritual path spanning foundational 48-day Mandalas, Rudraksha Diksha tiers (8, 11, 14 Mukhi), and advanced Charana initiations with approval gates.", GOLD_BRIGHT)
    ]

    card_h = Inches(1.15)
    gap_y = Inches(0.16)
    start_y = Inches(1.6)

    for i, (title, route, desc, col) in enumerate(pillars_info):
        cy = start_y + i * (card_h + gap_y)
        add_card(slide2, Inches(5.1), cy, Inches(7.433), card_h, bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
        tx_c = slide2.shapes.add_textbox(Inches(5.3), cy + Inches(0.1), Inches(7.0), card_h - Inches(0.2))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True

        p_t = tf_c.paragraphs[0]
        p_t.text = f"{title}   "
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        p_r = tf_c.add_paragraph()
        p_r.text = f"URL: {route}  |  {desc}"
        p_r.font.name = FONT_BODY
        p_r.font.size = Pt(10)
        p_r.font.color.rgb = TEXT_LIGHT

    slide2.notes_slide.notes_text_frame.text = (
        "Slide 2 establishes the high-level purpose of the super application: to unify the official website, "
        "the Q&A knowledge base, the Sahasralinga image codex, and the structured sadhana progression engine."
    )

    # =========================================================================
    # SLIDE 3: HIGH-LEVEL 4-TIER TECHNICAL ARCHITECTURE FLOW
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "System Architecture: 4-Tier Enterprise Flow", "Technical Foundations")
    add_footer(slide3, 3)

    tiers = [
        ("TIER 1: CLIENT PRESENTATION LAYER", [
            ("Unified Header & Chrome", "Glowing Devanagari ॐ, direct portal switches, audio bell chime"),
            ("Design Token System", "Void background (#07060a), radiant gold leaf (#d4a64a), ebony parchment"),
            ("Dynamic Interactive Components", "Canvas particle embers, ⌘K command modal, logarithmic spiral canvas"),
            ("Responsive Typography", "Cormorant Garamond, EB Garamond, Tiro Devanagari Sanskrit")
        ], GOLD_PRIMARY),
        ("TIER 2: FLASK CORE & ROUTING ENGINE", [
            ("Application Controller (app.py)", "Central WSGI entrypoint, route handlers, error handlers (404/500)"),
            ("Authentication Blueprint (auth.py)", "Flask-Login sessions, WTForms validation, registration vetting queue"),
            ("Security & Middleware", "CSRF protection, Werkzeug password hashing, role-based decorators"),
            ("Read-Only Sub-App Proxy", "Dynamic streaming proxy to local filesystem without altering original files")
        ], GOLD_BRIGHT),
        ("TIER 3: BUSINESS LOGIC & SERVICES", [
            ("Sādhana Stage Gatekeeper", "Calculates stage eligibility, duration timestamps, and permission checks"),
            ("Live Lunar Ashtami Engine", "Real-time calculation of Krishna Paksha Ashtami dates & night timings"),
            ("Real-Time Chat Dispatcher", "In-memory & DB messaging queue between Seekers and Guruji/Admin"),
            ("Google Sheets Sync Engine", "Bidirectional donation synchronization with cloud service accounts")
        ], GOLD_PRIMARY),
        ("TIER 4: DATA & PERSISTENCE LAYER", [
            ("Relational DB (SQLAlchemy)", "User, StageAccessRequest, MandalaRegistration, Chat, Donations"),
            ("Sacred Q&A Archive (Bhairva)", "257+ verified Q&A entries loaded into fast memory indices"),
            ("Media & Codex Storage", "324 high-res Shiva Lingam photographs, transcripts & metadata"),
            ("Excel Reporting (openpyxl)", "Automated seeker progression exports with embedded statistical charts")
        ], GOLD_BRIGHT),
    ]

    tier_w = Inches(2.78)
    tier_gap = Inches(0.2)
    start_x = Inches(0.8)

    for i, (tier_title, items, col) in enumerate(tiers):
        tx = start_x + i * (tier_w + tier_gap)
        add_card(slide3, tx, Inches(1.6), tier_w, Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)

        # Header banner inside card
        header_banner = add_card(slide3, tx, Inches(1.6), tier_w, Inches(0.85), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
        tx_h = slide3.shapes.add_textbox(tx + Inches(0.1), Inches(1.7), tier_w - Inches(0.2), Inches(0.7))
        p_th = tx_h.text_frame.paragraphs[0]
        p_th.text = tier_title
        p_th.font.name = FONT_HEADING
        p_th.font.size = Pt(11)
        p_th.font.bold = True
        p_th.font.color.rgb = col
        p_th.alignment = PP_ALIGN.CENTER

        # Items
        tx_items = slide3.shapes.add_textbox(tx + Inches(0.15), Inches(2.55), tier_w - Inches(0.3), Inches(4.0))
        tf_items = tx_items.text_frame
        tf_items.word_wrap = True

        for j, (item_h, item_d) in enumerate(items):
            p_ih = tf_items.paragraphs[0] if j == 0 else tf_items.add_paragraph()
            p_ih.text = f"▸ {item_h}"
            p_ih.font.name = FONT_BODY
            p_ih.font.size = Pt(10)
            p_ih.font.bold = True
            p_ih.font.color.rgb = GOLD_PRIMARY

            p_id = tf_items.add_paragraph()
            p_id.text = f"   {item_d}\n"
            p_id.font.name = FONT_BODY
            p_id.font.size = Pt(9)
            p_id.font.color.rgb = TEXT_LIGHT

    slide3.notes_slide.notes_text_frame.text = (
        "Slide 3 diagrams the technical architecture into four tiers: Presentation, Flask Routing, "
        "Domain Services, and Data Persistence. Highlight the zero-modification guarantee where local scraped "
        "data is proxied directly in read-only mode."
    )

    # =========================================================================
    # SLIDE 4: PUBLIC SEEKER DISCOVERY & KNOWLEDGE EXPLORATION FLOW
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Public Seeker Discovery & Knowledge Exploration Flow", "User Journey Flow 1")
    add_footer(slide4, 4)

    steps_flow1 = [
        ("Step 1: Sanctuary Entry", "Route: /", "Seeker lands on official homepage. Live Ashtami ribbon displays real-time countdown to next Krishna Paksha Ashtami. Hero invocation & ambient audio chime available.", "Entry Point"),
        ("Step 2: Sacred Discourses", "Routes: /guru-bhairava, /youtube", "Seeker views Daiva Anugraha YouTube videos, lineage discourses on Maa Kamakhya, Guru Bhairava, and the Avadhuta Siddha tradition.", "Media & Discourses"),
        ("Step 3: Codex Exploration", "Routes: /jnana-samvada, /bhairav-loka", "Deep dive into 257+ Q&As with instant ⌘K search and the interactive 324 Shiva Lingam Sahasralinga celestial spiral.", "Knowledge Archives"),
        ("Step 4: Sādhana Path Guidance", "Routes: /sadhana-paddhati, /devi-padathi", "Explores introductory 3-step discipline (Inner Dawn, Divine Deepening, Sacred Union) and Devi Kamakhya 33-99 day Mandalas.", "Spiritual Discipline"),
        ("Step 5: Registration & Commitment", "Routes: /auth/register, /mandala-sadhana", "Seeker formally commits: applies for gated seeker account or submits Mandala pledge with geo-location & vow declaration.", "Commitment Gateway")
    ]

    step_w = Inches(2.2)
    step_gap = Inches(0.18)
    start_x = Inches(0.8)

    for i, (s_title, s_route, s_desc, s_badge) in enumerate(steps_flow1):
        sx = start_x + i * (step_w + step_gap)
        
        # Step Container
        card = add_card(slide4, sx, Inches(1.8), step_w, Inches(4.8), bg_color=CARD_BG, border_color=CARD_BORDER)

        # Step Number Badge
        badge = add_card(slide4, sx + Inches(0.2), Inches(1.5), step_w - Inches(0.4), Inches(0.5), bg_color=CRIMSON_BG, border_color=CRIMSON)
        tx_b = slide4.shapes.add_textbox(sx + Inches(0.2), Inches(1.55), step_w - Inches(0.4), Inches(0.4))
        p_b = tx_b.text_frame.paragraphs[0]
        p_b.text = f"PHASE {i+1}"
        p_b.font.name = FONT_MONO
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = GOLD_BRIGHT
        p_b.alignment = PP_ALIGN.CENTER

        # Content
        tx_c = slide4.shapes.add_textbox(sx + Inches(0.12), Inches(2.15), step_w - Inches(0.24), Inches(4.3))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True

        p_t = tf_c.paragraphs[0]
        p_t.text = s_title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = GOLD_PRIMARY

        p_r = tf_c.add_paragraph()
        p_r.text = s_route
        p_r.font.name = FONT_MONO
        p_r.font.size = Pt(9)
        p_r.font.color.rgb = GOLD_DARK

        p_d = tf_c.add_paragraph()
        p_d.text = f"\n{s_desc}"
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_LIGHT

    # Flow arrows indicator at bottom
    flow_bar = add_card(slide4, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.5), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
    tx_fb = slide4.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.4))
    p_fb = tx_fb.text_frame.paragraphs[0]
    p_fb.text = "FLOW TRAJECTORY:   Public Sanctuary (/)  ➔  Knowledge Q&A & Codex  ➔  Sādhana Paddhati Guidance  ➔  Formal Seeker Registration"
    p_fb.font.name = FONT_MONO
    p_fb.font.size = Pt(10)
    p_fb.font.bold = True
    p_fb.font.color.rgb = GOLD_PRIMARY

    slide4.notes_slide.notes_text_frame.text = (
        "Slide 4 illustrates the public seeker discovery funnel. An unauthenticated visitor lands on the homepage, "
        "interacts with the Ashtami countdown, reads lineage wisdom, browses the Q&A and Sahasralinga codices, "
        "and ultimately decides to register or commit to a Mandala."
    )

    # =========================================================================
    # SLIDE 5: JNĀNA SAMVĀDA & BHAIRAV LOKA CODEX FLOWS
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Deep Dive: Jnāna Samvāda & Bhairav Loka Codex Systems", "Codex Architecture")
    add_footer(slide5, 5)

    # Left Box: Jnana Samvada (Q&A Codex)
    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_js = slide5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.7))
    tf_js = tx_js.text_frame
    tf_js.word_wrap = True

    p = tf_js.paragraphs[0]
    p.text = "☸ Jnāna Samvāda (Bhairava Q&A Codex)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GOLD_PRIMARY

    p_body = tf_js.add_paragraph()
    p_body.text = (
        "Archival repository containing 257+ verified answers by Guruji, "
        "organizing sacred knowledge into structured folios.\n\n"
        "Technical & Data Architecture:\n"
        "• Data Source: Read-only ingestion from Bhairva dataset via /api/qna.\n"
        "• Thematic Folios: 6 Core Categories:\n"
        "   1. Mantra & Japa (Seed syllables, recitation counts)\n"
        "   2. Pūjā, Āratī & Rituals (Lamp offerings, mustard oil, timing)\n"
        "   3. Maṇḍala & Anuṣṭhāna (48/144-day rules, fasting)\n"
        "   4. Experiences in Sādhanā (Visions, energy shifts, dreams)\n"
        "   5. Advanced Topics (Kundalini, cremation ground symbolism)\n"
        "   6. Women & Sādhanā (Specific guidance & sacred timings)\n"
        "• Instant Search: Client-side fuzzy matching with universal ⌘K shortcut.\n"
        "• Backward Compatibility: Seamless HTTP 301 redirection from legacy /qna."
    )
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(10)
    p_body.font.color.rgb = TEXT_LIGHT

    # Right Box: Bhairav Loka Sahasralinga Codex
    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_bl = slide5.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.333), Inches(4.7))
    tf_bl = tx_bl.text_frame
    tf_bl.word_wrap = True

    p_bl = tf_bl.paragraphs[0]
    p_bl.text = "✦ Bhairav Loka (Sahasralinga 324 Codex)"
    p_bl.font.name = FONT_HEADING
    p_bl.font.size = Pt(16)
    p_bl.font.bold = True
    p_bl.font.color.rgb = GOLD_BRIGHT

    p_bl_body = tf_bl.add_paragraph()
    p_bl_body.text = (
        "An interactive visualization embodying 324 consecrated Shiva Lingams, "
        "blending mathematical spiral aesthetics with sacred darshan.\n\n"
        "Key Engineering Features:\n"
        "• Logarithmic Celestial Spiral: HTML5 Canvas rendering 324 coordinate points along a golden-ratio spiral.\n"
        "• View Transformation Engine: Seeker can seamlessly switch between:\n"
        "   - Spiral Canvas (Celestial orbital view with zoom & pan)\n"
        "   - Darshan Grid (Organized by Lingam sequence)\n"
        "   - Dense Matrix (High-speed metadata lookup)\n"
        "   - Detailed Cards View (Expanded discourse view)\n"
        "• Modal Darshan Lightbox: High-resolution imagery proxied from instagram_scrap with verbatim discourse transcripts & hashtags.\n"
        "• Zero Latency: Pre-indexed JSON caching via /api/posts and /api/stats."
    )
    p_bl_body.font.name = FONT_BODY
    p_bl_body.font.size = Pt(10)
    p_bl_body.font.color.rgb = TEXT_LIGHT

    slide5.notes_slide.notes_text_frame.text = (
        "Slide 5 provides a side-by-side technical breakdown of the two major codices: "
        "Jnana Samvada with 257+ Q&As and Bhairav Loka with 324 consecrated Shiva Lingams rendered on a dynamic spiral canvas."
    )

    # =========================================================================
    # SLIDE 6: SEEKER REGISTRATION, VETTING & ADMIN APPROVAL WORKFLOW
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Seeker Registration, Gated Vetting & Admin Approval Architecture", "User Journey Flow 2")
    add_footer(slide6, 6)

    # Top 5-Step Stepper Ribbon
    stepper_steps = [
        ("1. Intake", "/auth/register"),
        ("2. DB Commit", "is_approved=False"),
        ("3. Gate Lock", "can_login()=False"),
        ("4. Admin Desk", "/auth/admin/users"),
        ("5. Activated", "Mandala 1 Unlocked")
    ]
    step_w = Inches(2.2)
    step_gap = Inches(0.18)
    for s_idx, (s_label, s_route) in enumerate(stepper_steps):
        sx = Inches(0.8) + s_idx * (step_w + step_gap)
        add_card(slide6, sx, Inches(1.5), step_w, Inches(0.65), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
        tx_s = slide6.shapes.add_textbox(sx + Inches(0.08), Inches(1.53), step_w - Inches(0.16), Inches(0.58))
        tf_s = tx_s.text_frame
        tf_s.word_wrap = True
        p_s1 = tf_s.paragraphs[0]
        p_s1.text = s_label
        p_s1.font.name = FONT_HEADING
        p_s1.font.size = Pt(10.5)
        p_s1.font.bold = True
        p_s1.font.color.rgb = GOLD_PRIMARY
        p_s1.alignment = PP_ALIGN.CENTER

        p_s2 = tf_s.add_paragraph()
        p_s2.text = s_route
        p_s2.font.name = FONT_MONO
        p_s2.font.size = Pt(8.5)
        p_s2.font.color.rgb = TEXT_MUTED
        p_s2.alignment = PP_ALIGN.CENTER

    # 3-Tier Swimlane Architecture
    swimlanes = [
        ("1. Seeker Experience", "Frontend Aspirant Journey", [
            "Visits /auth/register and completes profile form.",
            "Submits mandatory Purpose Essay (20-1000 chars) detailing spiritual resolve.",
            "Redirected to /auth/login with queue confirmation notice.",
            "Pre-approval login attempts show protective status advisory.",
            "Post-approval login initiates 48-day Mandala 1 timer and unlocks personal Sādhana Path."
        ], GOLD_PRIMARY),
        ("2. Security & DB Gateway", "Automated Middleware Engine", [
            "WTForms validates email format, password strength & unique constraints.",
            "Password encrypted via PBKDF2:SHA256 with secure random salt.",
            "User committed with is_approved=False, is_active=True, role='user'.",
            "can_login() gate prevents unauthorized session initiation.",
            "inject_admin_notifications updates real-time pending applicant counter.",
            "Approval sets is_approved=True, records approved_at & sets mandala_1_started_at."
        ], GOLD_BRIGHT),
        ("3. Admin Operations Desk", "Guruji & Mentors Console", [
            "Top navigation bar alerts admin with glowing pending count badge.",
            "Admin accesses /auth/admin/users (redirects from /admin/users).",
            "Reads applicant's purpose modal & assesses spiritual experience tier.",
            "Tri-State Decision: Approve (unlocks Mandala 1), Reject, or Suspend.",
            "Audit trail recorded with approved_by ID and ISO timestamp.",
            "Admin can dispatch direct spiritual guidance via in-app mentorship chat."
        ], GOLD_PRIMARY)
    ]

    col_w = Inches(3.75)
    gap_x = Inches(0.24)
    start_x = Inches(0.8)

    for i, (title, sub, bullet_pts, col) in enumerate(swimlanes):
        cx = start_x + i * (col_w + gap_x)
        add_card(slide6, cx, Inches(2.3), col_w, Inches(3.85), bg_color=CARD_BG, border_color=CARD_BORDER)

        # Header Pill
        header_pill = add_card(slide6, cx, Inches(2.3), col_w, Inches(0.75), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
        tx_p = slide6.shapes.add_textbox(cx + Inches(0.1), Inches(2.35), col_w - Inches(0.2), Inches(0.65))
        tf_p = tx_p.text_frame
        tf_p.word_wrap = True
        p_pt = tf_p.paragraphs[0]
        p_pt.text = title
        p_pt.font.name = FONT_HEADING
        p_pt.font.size = Pt(12)
        p_pt.font.bold = True
        p_pt.font.color.rgb = col
        p_pt.alignment = PP_ALIGN.CENTER

        p_ps = tf_p.add_paragraph()
        p_ps.text = sub
        p_ps.font.name = FONT_BODY
        p_ps.font.size = Pt(9)
        p_ps.font.color.rgb = TEXT_MUTED
        p_ps.alignment = PP_ALIGN.CENTER

        # Body items
        tx_b = slide6.shapes.add_textbox(cx + Inches(0.12), Inches(3.15), col_w - Inches(0.24), Inches(2.9))
        tf_b = tx_b.text_frame
        tf_b.word_wrap = True

        for j, b_text in enumerate(bullet_pts):
            p_bullet = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            p_bullet.text = f"• {b_text}"
            p_bullet.font.name = FONT_BODY
            p_bullet.font.size = Pt(9.5)
            p_bullet.font.color.rgb = TEXT_LIGHT

    # Bottom Callout: Security & Spiritual Purity Gate
    callout = add_card(slide6, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.55), bg_color=CARD_BG_LIGHT, border_color=GOLD_PRIMARY)
    tx_co = slide6.shapes.add_textbox(Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.45))
    p_co = tx_co.text_frame.paragraphs[0]
    p_co.text = "✦ ZERO-FRICTION SPIRITUAL GATEKEEPING: From initial intake to cryptographic storage, human vetting, and stage clock initiation, this unified flow balances seamless onboarding with sacred lineage protection."
    p_co.font.name = FONT_BODY
    p_co.font.size = Pt(10)
    p_co.font.bold = True
    p_co.font.color.rgb = GOLD_BRIGHT

    slide6.notes_slide.notes_text_frame.text = (
        "Slide 6 walks through the complete end-to-end journey from seeker intake and cryptographic storage, "
        "to gatekeeper pending state, admin vetting at /auth/admin/users, and automatic Mandala 1 clock initiation."
    )

    # =========================================================================
    # SLIDE 7: SĀDHANA STAGE PROGRESSION & DIKSHA PROTOCOL
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Sādhana Stage Progression & Diksha Unlock Protocol", "Spiritual Hierarchy Flow")
    add_footer(slide7, 7)

    # Top Grid: The 9 Progressive Stages
    stages_data = [
        ("Stage 1: Mandala 1", "48-Day Foundation", "Prathama Mandala. Basic lamp offerings, Bhairava mantra recitation. Auto-unlocked upon approval."),
        ("Stage 2: Mandala 2", "48-Day Deepening", "Dvitiya Mandala. Advanced japa discipline, dietary tapasya. Unlocked upon Stage 1 completion."),
        ("Stage 3: Mandala 3", "48-Day Mastery", "Tritiya Mandala. Intensive 144-day cumulative discipline, deeper meditation timings."),
        ("Stage 4: 8-Mukhi Diksha", "Ashta Murti", "Rudraksha Diksha for overcoming planetary obstacles & establishing fearless mental ground."),
        ("Stage 5: 11-Mukhi Diksha", "Ekadasha Rudra", "Rudraksha Diksha invoking 11 forms of Rudra. Protection, higher discrimination, japa intensification."),
        ("Stage 6: 14-Mukhi Diksha", "Deva Mani", "Awakening Ajna Chakra, direct perception of Shiva Tattva. Advanced seeker milestone."),
        ("Stage 7: Pratham Charana", "Diksha Inception", "Introduction to high-esoteric Rudraksha Diksha with consecrated rituals & sankalpas."),
        ("Stage 8: Dutiya Charana", "Second Phase", "Esoteric immersion for mature sādhanā practitioners under direct guidance of Guruji."),
        ("Stage 9: Tritiya Charana", "Supreme Realization", "Culminating Siddha transmission and total immersion into Sarvam Bhairava Swaroopam.")
    ]

    col_w = Inches(3.75)
    row_h = Inches(1.1)
    gap_x = Inches(0.24)
    gap_y = Inches(0.12)

    for idx, (st_name, st_sub, st_desc) in enumerate(stages_data):
        row = idx // 3
        col = idx % 3
        cx = Inches(0.8) + col * (col_w + gap_x)
        cy = Inches(1.6) + row * (row_h + gap_y)

        add_card(slide7, cx, cy, col_w, row_h, bg_color=CARD_BG, border_color=CARD_BORDER)
        tx = slide7.shapes.add_textbox(cx + Inches(0.1), cy + Inches(0.08), col_w - Inches(0.2), row_h - Inches(0.16))
        tf = tx.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = f"{st_name} "
        p_h.font.name = FONT_HEADING
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = GOLD_PRIMARY

        p_s = tf.add_paragraph()
        p_s.text = f"[{st_sub}]  {st_desc}"
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(9)
        p_s.font.color.rgb = TEXT_LIGHT

    # Bottom Box: Stage Request & Approval Protocol
    protocol_box = add_card(slide7, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.5), bg_color=CARD_BG_LIGHT, border_color=GOLD_PRIMARY)
    tx_pr = slide7.shapes.add_textbox(Inches(1.0), Inches(5.45), Inches(11.3), Inches(1.3))
    tf_pr = tx_pr.text_frame
    tf_pr.word_wrap = True

    p = tf_pr.paragraphs[0]
    p.text = "✦ STAGE UNLOCK STATE MACHINE & VERIFICATION CYCLE:"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = GOLD_BRIGHT

    p_p = tf_pr.add_paragraph()
    p_p.text = (
        "1. Seeker Finishes Current Stage  ➔  Clicks 'Request Next Stage Access' in Profile or Stage Portal.\n"
        "2. StageAccessRequest Created  ➔  Logged with user_id, target stage_number, requested_at, status='pending'.\n"
        "3. Admin Review & Verification  ➔  Admin inspects duration spent (get_stage_duration_days), chat questions & commitment.\n"
        "4. Grant & State Transition  ➔  Admin approves; sets stage_access=True, complete_stage() stamps completion date, and start_stage() initiates duration timer."
    )
    p_p.font.name = FONT_BODY
    p_p.font.size = Pt(10)
    p_p.font.color.rgb = TEXT_LIGHT

    slide7.notes_slide.notes_text_frame.text = (
        "Slide 7 explains the structured 9-stage progression. Detail the mechanics of StageAccessRequest, "
        "where seekers must request access to subsequent stages, and admins verify duration and spiritual commitment before unlocking."
    )

    # =========================================================================
    # SLIDE 8: MANDALA SĀDHANA REGISTRATION & PLEDGE FLOW
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Mandala Sādhana Commitment & Pledge Flow", "Dedicated Spiritual Vows")
    add_footer(slide8, 8)

    # Left: Seeker Vow Submission Flow
    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_m1 = slide8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.7))
    tf_m1 = tx_m1.text_frame
    tf_m1.word_wrap = True

    p_m1 = tf_m1.paragraphs[0]
    p_m1.text = "Seeker Pledge Submission (/mandala-sadhana)"
    p_m1.font.name = FONT_HEADING
    p_m1.font.size = Pt(15)
    p_m1.font.bold = True
    p_m1.font.color.rgb = GOLD_PRIMARY

    p_m1_desc = tf_m1.add_paragraph()
    p_m1_desc.text = (
        "A formal vow portal where seekers commit to unbroken discipline for 48 or 144 days.\n\n"
        "Captured Data Points:\n"
        "• Full Name & Geo-Location (City, Country)\n"
        "• Email address for response verification\n"
        "• 48-Day Mandala Commitment (Mandatory boolean vow)\n"
        "• 144-Day Extended Commitment (Yes / No / Not Yet Ready)\n"
        "• Sādhana Category (Bhairava Japa, Batuk Bhairava, Devi Kamakhya)\n"
        "• Declared Start Date (sadhana_start_date)\n"
        "• Deep Personal Commitment Statement (commitment_text)\n"
        "• Option to receive confirmation copy (send_copy)\n\n"
        "Validation Engine:\n"
        "Handled via Flask route /mandala-sadhana and REST endpoint /api/mandala-sadhana with strict integrity checks."
    )
    p_m1_desc.font.name = FONT_BODY
    p_m1_desc.font.size = Pt(10)
    p_m1_desc.font.color.rgb = TEXT_LIGHT

    # Right: Admin Roster & Export Flow
    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_m2 = slide8.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.333), Inches(4.7))
    tf_m2 = tx_m2.text_frame
    tf_m2.word_wrap = True

    p_m2 = tf_m2.paragraphs[0]
    p_m2.text = "Admin Pledge Governance (/admin/mandala-sadhana)"
    p_m2.font.name = FONT_HEADING
    p_m2.font.size = Pt(15)
    p_m2.font.bold = True
    p_m2.font.color.rgb = GOLD_BRIGHT

    p_m2_desc = tf_m2.add_paragraph()
    p_m2_desc.text = (
        "The administrative desk oversees the entire global collective of active mandalas.\n\n"
        "Operational Capabilities:\n"
        "• Central Registration Roster: Real-time table displaying all seekers, start dates, and vows.\n"
        "• Detailed Pledge View (/admin/mandala-sadhana/<id>): Inspect seeker vows and contact info.\n"
        "• Search & Filter: Filter by commitment duration (48 vs 144 days) and sadhana type.\n"
        "• Data Export Engine (/admin/mandala-sadhana/export): Instant CSV / Excel export for Guruji's ritual sankalpa rosters.\n"
        "• Cross-Reference with User Accounts: Links anonymous or external pledges with verified user profiles when emails match."
    )
    p_m2_desc.font.name = FONT_BODY
    p_m2_desc.font.size = Pt(10)
    p_m2_desc.font.color.rgb = TEXT_LIGHT

    slide8.notes_slide.notes_text_frame.text = (
        "Slide 8 covers the dedicated Mandala Sadhana vow submission and management flow. "
        "Explain how the application records 48-day and 144-day vows and allows administrators to export the roster."
    )

    # =========================================================================
    # SLIDE 9: REAL-TIME SEEKER-TO-ADMIN GUIDANCE CHAT FLOW
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Real-Time Seeker & Admin Guidance Chat Flow", "Spiritual Support Architecture")
    add_footer(slide9, 9)

    # 3-Column Chat Architecture
    chat_cols = [
        ("1. Seeker Chat Interface", [
            ("Persistent Floating Widget", "Embedded on all authenticated seeker pages; expandable floating icon."),
            ("Inquiry Submission", "POST /api/chat/send sends message text with seeker authentication."),
            ("Chat History Polling", "GET /api/chat/history retrieves chronological messages between seeker and admin."),
            ("Unread Indicator", "Real-time badge indicating pending spiritual guidance responses from Guruji.")
        ], GOLD_PRIMARY),
        ("2. Backend Message Daemon", [
            ("ChatMessage Model", "Persists sender_id, recipient_id, message text, timestamp, and is_from_admin flag."),
            ("Routing Intelligence", "When seeker sends, recipient defaults to admin pool; when admin replies, user is targeted."),
            ("Read-State Synchronization", "POST /api/chat/mark-read updates is_read flag upon opening conversation drawer."),
            ("Zero WebSocket Overhead", "High-efficiency lightweight REST polling designed for serverless reliability.")
        ], GOLD_BRIGHT),
        ("3. Admin Chat Operations Desk", [
            ("Unified Console (/admin/chat)", "Full-screen multi-user management dashboard for administrative guidance."),
            ("Seeker Roster & Unread Badges", "GET /api/admin/chat/users lists all seekers with unread counts & last active dates."),
            ("Real-Time Message View", "GET /api/admin/chat/<user_id> loads focused conversation history."),
            ("Instant Spiritual Response", "POST to /api/chat/send with is_from_admin=True dispatches immediate guidance.")
        ], GOLD_PRIMARY)
    ]

    col_w = Inches(3.75)
    gap_x = Inches(0.24)
    start_x = Inches(0.8)

    for i, (title, items, col) in enumerate(chat_cols):
        cx = start_x + i * (col_w + gap_x)
        add_card(slide9, cx, Inches(1.6), col_w, Inches(4.5), bg_color=CARD_BG, border_color=CARD_BORDER)

        # Header
        hdr = add_card(slide9, cx, Inches(1.6), col_w, Inches(0.7), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
        tx_h = slide9.shapes.add_textbox(cx + Inches(0.1), Inches(1.65), col_w - Inches(0.2), Inches(0.6))
        p_th = tx_h.text_frame.paragraphs[0]
        p_th.text = title
        p_th.font.name = FONT_HEADING
        p_th.font.size = Pt(12)
        p_th.font.bold = True
        p_th.font.color.rgb = col
        p_th.alignment = PP_ALIGN.CENTER

        # Body items
        tx_b = slide9.shapes.add_textbox(cx + Inches(0.15), Inches(2.4), col_w - Inches(0.3), Inches(3.6))
        tf_b = tx_b.text_frame
        tf_b.word_wrap = True

        for j, (h_txt, d_txt) in enumerate(items):
            p_ih = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            p_ih.text = f"▸ {h_txt}"
            p_ih.font.name = FONT_BODY
            p_ih.font.size = Pt(10)
            p_ih.font.bold = True
            p_ih.font.color.rgb = GOLD_PRIMARY

            p_id = tf_b.add_paragraph()
            p_id.text = f"   {d_txt}\n"
            p_id.font.name = FONT_BODY
            p_id.font.size = Pt(9)
            p_id.font.color.rgb = TEXT_LIGHT

    # Bottom Diagram
    flow_dia = add_card(slide9, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.55), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
    tx_dia = slide9.shapes.add_textbox(Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.45))
    p_dia = tx_dia.text_frame.paragraphs[0]
    p_dia.text = "CHAT FLOW:   Seeker submits inquiry  ➔  REST endpoint stores in DB  ➔  Admin Console notifies  ➔  Admin replies  ➔  Seeker notified"
    p_dia.font.name = FONT_MONO
    p_dia.font.size = Pt(10)
    p_dia.font.bold = True
    p_dia.font.color.rgb = GOLD_BRIGHT

    slide9.notes_slide.notes_text_frame.text = (
        "Slide 9 demonstrates the custom chat architecture. "
        "Because spiritual guidance during intense mandalas requires personal instruction, "
        "the built-in chat system connects authenticated seekers directly with administrators without relying on third-party SaaS."
    )

    # =========================================================================
    # SLIDE 10: ADMINISTRATIVE GOVERNANCE & BHIKSHA (DONATIONS) SYNC FLOW
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Administrative Governance & Bhiksha (Donations) Sync Flow", "Admin & Financial Operations")
    add_footer(slide10, 10)

    # Left: Seeker Management & Excel Analytics
    add_card(slide10, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_a1 = slide10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.7))
    tf_a1 = tx_a1.text_frame
    tf_a1.word_wrap = True

    p = tf_a1.paragraphs[0]
    p.text = "Seeker Governance & Excel Analytics"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GOLD_PRIMARY

    p_body = tf_a1.add_paragraph()
    p_body.text = (
        "Central command center for seeker management (/auth/users):\n\n"
        "• Granular Stage Permission Toggles:\n"
        "   Direct admin toggles for Mandalas 1-3, Rudrakshas (8, 11, 14 Mukhi), and Charana Dikshas.\n"
        "• Seeker Lifecycle Auditing:\n"
        "   Tracks created_at, approved_at, last_active, and current stage duration.\n"
        "• Advanced User Search & Filter:\n"
        "   Filter by approval status, practice level, stage number, or referral source.\n"
        "• Automated Excel Analytics Engine (openpyxl):\n"
        "   Generates multi-sheet workbook with seeker directory, stage distribution metrics, and embedded bar charts."
    )
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(10)
    p_body.font.color.rgb = TEXT_LIGHT

    # Right: Bhiksha & Google Sheets Integration
    add_card(slide10, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_a2 = slide10.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.333), Inches(4.7))
    tf_a2 = tx_a2.text_frame
    tf_a2.word_wrap = True

    p_a2 = tf_a2.paragraphs[0]
    p_a2.text = "Bhiksha (Donations) & Google Sheets Sync"
    p_a2.font.name = FONT_HEADING
    p_a2.font.size = Pt(15)
    p_a2.font.bold = True
    p_a2.font.color.rgb = GOLD_BRIGHT

    p_a2_body = tf_a2.add_paragraph()
    p_a2_body.text = (
        "Transparent management of sacred offerings & financial contributions:\n\n"
        "• Donation Purpose Categories (DonationPurpose):\n"
        "   Annadanam, Temple Construction, Sādhana Kutir, Puja Samagri.\n"
        "• Offline Donation Verification Queue:\n"
        "   Records cash, UPI, bank transfers with transaction reference verification.\n"
        "• Google Sheets Cloud Sync Engine (google_sheets.py):\n"
        "   - Connects to Google Cloud Service Account.\n"
        "   - Automatic pulling & matching of donor records from worksheets.\n"
        "   - Tracks donor_id, worksheet name, amount, and payment date.\n"
        "   - Manual sync triggers (/admin/sync-now) & background API sync."
    )
    p_a2_body.font.name = FONT_BODY
    p_a2_body.font.size = Pt(10)
    p_a2_body.font.color.rgb = TEXT_LIGHT

    slide10.notes_slide.notes_text_frame.text = (
        "Slide 10 details administrative governance and financial management. "
        "Highlight the openpyxl integration that generates Excel reports with embedded charts and "
        "the Google Sheets engine that syncs donation contributions."
    )

    # =========================================================================
    # SLIDE 11: DEVOPS, SECURITY & RELIABILITY ARCHITECTURE
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "DevOps, Deployment, Security & Reliability Architecture", "System Governance")
    add_footer(slide11, 11)

    cards_devops = [
        ("Zero-Modification Architecture", [
            "Local upstream repositories (Bhairva & instagram_scrap) remain 100% clean and untouched.",
            "Flask routes dynamically read and proxy media and JSON datasets in read-only mode.",
            "Guarantees that git pulls on upstream codices never encounter merge conflicts."
        ], GOLD_PRIMARY),
        ("Security & Access Control", [
            "Werkzeug PBKDF2:SHA256 password hashing with salt.",
            "Role-Based Access Control (RBAC): Admin vs Seeker decorators.",
            "CSRF protection on all forms via Flask-WTF.",
            "Gated login verification preventing unvetted seeker access."
        ], GOLD_BRIGHT),
        ("Automated Verification Suite", [
            "Comprehensive integration test suite in test_superapp.py.",
            "Verifies all 6 primary HTML routes, 4 API endpoints, and proxied media assets.",
            "Maintains 100% pass rate (200 OK across 17 checkpoints) before deployment."
        ], GOLD_PRIMARY),
        ("Vercel Serverless Deployment", [
            "Optimized for serverless WSGI execution via vercel.json.",
            "Static asset routing directly to static/ directory with edge caching.",
            "Graceful fallback handling with custom 404 and 500 temple error pages."
        ], GOLD_BRIGHT)
    ]

    card_w = Inches(5.7)
    card_h = Inches(2.4)
    gap_x = Inches(0.33)
    gap_y = Inches(0.25)

    for i, (title, points, col) in enumerate(cards_devops):
        row = i // 2
        col_idx = i % 2
        cx = Inches(0.8) + col_idx * (card_w + gap_x)
        cy = Inches(1.6) + row * (card_h + gap_y)

        add_card(slide11, cx, cy, card_w, card_h, bg_color=CARD_BG, border_color=CARD_BORDER)
        tx = slide11.shapes.add_textbox(cx + Inches(0.15), cy + Inches(0.12), card_w - Inches(0.3), card_h - Inches(0.24))
        tf = tx.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = FONT_HEADING
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = col

        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_BODY
            p_pt.font.size = Pt(9.5)
            p_pt.font.color.rgb = TEXT_LIGHT

    # Reliability Footer
    rel_box = add_card(slide11, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.5), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
    tx_rel = slide11.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.4))
    p_rel = tx_rel.text_frame.paragraphs[0]
    p_rel.text = "✦ RELIABILITY STATUS: 17/17 Integration Tests Passing  •  Zero Upstream Mod Friction  •  Production-Ready WSGI Architecture"
    p_rel.font.name = FONT_MONO
    p_rel.font.size = Pt(10)
    p_rel.font.bold = True
    p_rel.font.color.rgb = GOLD_PRIMARY

    slide11.notes_slide.notes_text_frame.text = (
        "Slide 11 highlights engineering quality and DevOps principles: "
        "the Zero-Modification guarantee, security practices, the automated test suite, and cloud readiness."
    )

    # =========================================================================
    # SLIDE 12: SYSTEM ENTITY RELATIONSHIP & DATA MODEL FLOW
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "Data Architecture: Core Entities & Relationships", "Entity-Relationship Flow")
    add_footer(slide12, 12)

    # 4 Data Model Cards
    models = [
        ("User Model", "Primary Seeker Entity", [
            "id (PK), username, email, password_hash",
            "full_name, phone, address, role ('admin'/'user')",
            "is_approved, is_active, approved_at, approved_by",
            "practice_level, purpose, profile_picture",
            "mandala_1/2/3_access (Boolean flags)",
            "rudraksha_8/11/14_mukhi_access",
            "pratham/dutiya/tritiya_charana_access",
            "devi_mandala_1/2/3_access",
            "Stage start & completion timestamps"
        ], GOLD_PRIMARY),
        ("StageAccessRequest", "Progression Requests", [
            "id (PK), user_id (FK -> User)",
            "stage_number (1 to 9, or 101-103)",
            "status ('pending', 'approved', 'rejected')",
            "requested_at (DateTime)",
            "reviewed_at (DateTime)",
            "reviewed_by (FK -> User)",
            "notes (Text remarks by admin)"
        ], GOLD_BRIGHT),
        ("ChatMessage", "Spiritual Guidance Chat", [
            "id (PK), sender_id (FK -> User)",
            "recipient_id (FK -> User, optional)",
            "message (Text)",
            "timestamp (DateTime)",
            "is_read (Boolean)",
            "is_from_admin (Boolean)",
            "to_dict() JSON serializer"
        ], GOLD_PRIMARY),
        ("Donation & Mandala", "Offerings & Pledges", [
            "DonationPurpose: name, desc, is_active",
            "OfflineDonation: donor_name, amount,",
            "  currency, purpose_id, method, is_verified,",
            "  donor_id, worksheet (Google Sheets sync)",
            "MandalaSadhanaRegistration:",
            "  email, full_name, 48/144_commitment,",
            "  commitment_text, sadhana_start_date"
        ], GOLD_BRIGHT)
    ]

    card_w = Inches(2.78)
    gap_x = Inches(0.2)
    start_x = Inches(0.8)

    for i, (m_name, m_sub, attrs, col) in enumerate(models):
        cx = start_x + i * (card_w + gap_x)
        add_card(slide12, cx, Inches(1.6), card_w, Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)

        # Header
        hdr = add_card(slide12, cx, Inches(1.6), card_w, Inches(0.75), bg_color=CARD_BG_LIGHT, border_color=CARD_BORDER)
        tx_h = slide12.shapes.add_textbox(cx + Inches(0.1), Inches(1.65), card_w - Inches(0.2), Inches(0.65))
        p_th = tx_h.text_frame.paragraphs[0]
        p_th.text = m_name
        p_th.font.name = FONT_HEADING
        p_th.font.size = Pt(12)
        p_th.font.bold = True
        p_th.font.color.rgb = col
        p_th.alignment = PP_ALIGN.CENTER

        p_ts = tx_h.text_frame.add_paragraph()
        p_ts.text = m_sub
        p_ts.font.name = FONT_BODY
        p_ts.font.size = Pt(9)
        p_ts.font.color.rgb = TEXT_MUTED
        p_ts.alignment = PP_ALIGN.CENTER

        # Attributes
        tx_b = slide12.shapes.add_textbox(cx + Inches(0.12), Inches(2.45), card_w - Inches(0.24), Inches(4.1))
        tf_b = tx_b.text_frame
        tf_b.word_wrap = True

        for j, attr in enumerate(attrs):
            p_a = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            p_a.text = f"• {attr}"
            p_a.font.name = FONT_MONO if '(' in attr else FONT_BODY
            p_a.font.size = Pt(9.5)
            p_a.font.color.rgb = TEXT_LIGHT

    slide12.notes_slide.notes_text_frame.text = (
        "Slide 12 maps the core database schema: User, StageAccessRequest, ChatMessage, "
        "Donation models, and MandalaSadhanaRegistration."
    )

    # =========================================================================
    # SLIDE 13: STRATEGIC SUMMARY & FUTURE ROADMAP
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13)
    add_header(slide13, "Executive Summary & Future Evolution Roadmap", "Platform Synthesis")
    add_footer(slide13, 13)

    # Left: Key Accomplishments
    add_card(slide13, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_s1 = slide13.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.7))
    tf_s1 = tx_s1.text_frame
    tf_s1.word_wrap = True

    p = tf_s1.paragraphs[0]
    p.text = "Accomplishments & Current Metrics"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GOLD_PRIMARY

    p_body = tf_s1.add_paragraph()
    p_body.text = (
        "✦ Unified Sacred Experience:\n"
        "   Consolidated bhairavaanugraha.com, Jnāna Samvāda (257+ Q&As), and Sahasralinga 324 into a single portal.\n\n"
        "✦ Complete Sādhana Gatekeeping:\n"
        "   9-stage hierarchical spiritual progression engine ensuring proper seeker vetting and discipline.\n\n"
        "✦ Robust Administration & Sync:\n"
        "   User approval desk, real-time guidance chat, automated Excel exports with charts, and Google Sheets donation sync.\n\n"
        "✦ 100% Verified Zero-Modification Integration:\n"
        "   Upstream datasets remain pristine; verified across 17 automated integration checkpoints."
    )
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = TEXT_LIGHT

    # Right: Future Roadmap
    add_card(slide13, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_s2 = slide13.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.333), Inches(4.7))
    tf_s2 = tx_s2.text_frame
    tf_s2.word_wrap = True

    p_s2 = tf_s2.paragraphs[0]
    p_s2.text = "Strategic Future Roadmap"
    p_s2.font.name = FONT_HEADING
    p_s2.font.size = Pt(16)
    p_s2.font.bold = True
    p_s2.font.color.rgb = GOLD_BRIGHT

    p_s2_body = tf_s2.add_paragraph()
    p_s2_body.text = (
        "1. Progressive Web App (PWA) & Native Mobile App:\n"
        "   Offline access to mantra texts, stotras, and daily japa counters on Android and iOS.\n\n"
        "2. Automated Ashtami & Rahu Kalam Push Notifications:\n"
        "   Web push & mobile notifications alerting registered seekers 2 hours prior to Krishna Paksha Ashtami night portals.\n\n"
        "3. Live High-Fidelity Audio Chanting Stream:\n"
        "   Ambient temple stotras, bell resonances, and guided audio sādhanā tracks with synchronized Devanagari lyrics.\n\n"
        "4. Multilingual Codex Expansion:\n"
        "   Full multilingual localization for Hindi, Kannada, Telugu, and Tamil seekers."
    )
    p_s2_body.font.name = FONT_BODY
    p_s2_body.font.size = Pt(10.5)
    p_s2_body.font.color.rgb = TEXT_LIGHT

    slide13.notes_slide.notes_text_frame.text = (
        "Slide 13 concludes the presentation by summarizing the application's accomplishments and outlining "
        "the future roadmap, including PWA, automated push notifications, audio chanting streams, and multilingual localization."
    )

    # Save presentation
    output_filename = "Bhairava_Anugraha_Flow_Presentation.pptx"
    output_path = os.path.join(SCRIPT_DIR, output_filename)
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")
    return output_path


if __name__ == '__main__':
    build_presentation()
